"""Étape 3 — Index : sommaire des pièces, ordre des pages et ordre
chronologique, dans `01_index.xlsx`."""

from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from rich.console import Console

COLONNES = ["Type", "Page début", "Page fin", "Cote", "Date apparente", "Service rédacteur", "Confiance", "Statut"]


def _cle_tri_date(date_apparente: str | None, heure_apparente: str | None) -> tuple[int, datetime]:
    if not date_apparente:
        return (1, datetime.max)
    try:
        j, m, a = date_apparente.split("/")
        heure, minute = 0, 0
        if heure_apparente:
            hh, mm = heure_apparente.split("h")
            heure, minute = int(hh), int(mm)
        return (0, datetime(int(a), int(m), int(j), heure, minute))
    except ValueError:
        return (1, datetime.max)


def _ecrire_feuille(ws: Worksheet, pieces: list[sqlite3.Row]) -> None:
    ws.append(COLONNES)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for p in pieces:
        ws.append(
            [
                p["type"],
                p["page_debut"],
                p["page_fin"],
                p["cote"] or "NON TROUVÉ",
                p["date_apparente"] or "NON TROUVÉ",
                p["service_redacteur"] or "NON TROUVÉ",
                round(p["confiance"], 2),
                p["statut_revision"],
            ]
        )
    for colonne in ws.columns:
        largeur = max(len(str(c.value)) for c in colonne) + 2
        ws.column_dimensions[colonne[0].column_letter].width = min(largeur, 60)


def construire_index(db: sqlite3.Connection, affaire_dir: Path, console: Console) -> None:
    pieces = db.execute("SELECT * FROM pieces ORDER BY page_debut").fetchall()
    if not pieces:
        console.print("  [index] aucune pièce en base — lance d'abord `depouille classify`.")
        return

    pieces_chrono = sorted(
        pieces, key=lambda p: (_cle_tri_date(p["date_apparente"], p["heure_apparente"]), p["page_debut"])
    )

    wb = Workbook()
    feuille_pages = wb.active
    feuille_pages.title = "Ordre des pages"
    _ecrire_feuille(feuille_pages, pieces)

    feuille_chrono = wb.create_sheet("Ordre chronologique")
    _ecrire_feuille(feuille_chrono, pieces_chrono)

    dossier_out = affaire_dir / "out"
    dossier_out.mkdir(parents=True, exist_ok=True)
    chemin = dossier_out / "01_index.xlsx"
    wb.save(chemin)

    console.print(f"  [index] {len(pieces)} pièces -> {chemin}")
